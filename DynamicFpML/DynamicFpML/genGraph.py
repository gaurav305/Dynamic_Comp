from openpyxl import Workbook

from openpyxl.chart import (
    PieChart,
    ProjectedPieChart,
    Reference
)
from openpyxl.chart.series import DataPoint
wb = Workbook()
values_list = [['Trade ID', 'Pass', 'Fail', 'Total'],
               ['ts4-SFXFLXTDParabellum-fx-1509178000000091', 20, 7, 27],
               ['Test_JPM1', 43, 7, 50],
               ['Test_BBG1', 10, 7, 17]]
ws = wb.active

for row in values_list:
    ws.append(row)



# labels = Reference(ws, min_col=1, min_row=2, max_row=5)
# data = Reference(ws, min_col=2, min_row=1, max_row=5)
row = 2
chart = 1
for data in range(len(values_list)-1):
    pie = PieChart()

    labels_1 = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=1)
    data_p = Reference(ws, min_col=2, max_col=3, min_row=row, max_row=row)


    pie.add_data(data_p, titles_from_data=True)
    pie.set_categories(labels_1)
    pie.title = values_list[row-1][0]
    ws.add_chart(pie, "F{}".format(chart))
    row += 1
    chart = row*14

"""for x in values_list:

    rr = i+1
    labels_1 = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=1)
    # data_p = Reference(ws, min_col=2, max_col=4, min_row=2, max_row=2)
    data_p = Reference(ws, min_col=2, max_col=3, min_row=rr, max_row=rr)

    pie.add_data(data_p, titles_from_data=True)
    pie.set_categories(labels_1)
    pie.title = "{}".format(x[0])
    ws.add_chart(pie, "F{}".format(i*19))
    i += 1"""

wb.save("C:\Users\chayan.mazumder\Desktop\CompareResult2\pie3.xlsx")
